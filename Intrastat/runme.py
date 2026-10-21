#!/usr/bin/env python3
"""Intrastat pipeline for the Copilot agent.

Commands
  run    Step 1. Build "Intrastat_YYYY-MM.xlsm", the review workbook + JSON and
         the updated LEV-HSKoder from the monthly inputs. Also writes a run
         manifest ("Intrastat_run_YYYY-MM.json") so step 2 can rebuild
         everything from exactly the same inputs.
  codes  Step 2 helper. Lists valid 8-digit codes from the template's Opslag
         sheet (filter by --prefix and/or --search) so only codes the template
         accepts are chosen.
  apply  Step 2. Reads a decisions JSON with commodity codes for REVIEW
         invoices, validates every code against Opslag and rebuilds all outputs
         from the ORIGINAL inputs with the accepted codes applied.
         All valid codes are applied immediately, whatever their confidence;
         confidence is recorded so uncertain codes can be checked afterwards.
         --hold-low keeps low-confidence codes (and medium "service"
         decisions) in review instead.

Running without a command behaves like "run" (backwards compatible).

Inputs for "run" (auto-discovered in --data-dir when not given):
  --source           Monthly Fusion Intrastat export (.xlsx/.xlsm)
  --reference        Intrastat template (.xlsm, name contains "skabelon")
  --period           YYYY-MM (inferred from the source when omitted)
  --hskoder          LEV-HSKoder.xlsx (supplier history; highest count wins,
                     first listed code wins a tie)
  --invoice-details  Fakturaer YYYY-MM.xlsx (supplementary quantities and
                     invoice-line descriptions for the review)
  --overrides        CSV/XLSX with explicit invoice corrections
  --output-dir       Destination folder (default: <data-dir>/output)

Decisions file for "apply":
  {"period": "YYYY-MM",
   "decisions": [{"invoice_number": "...", "commodity_code": "84713000",
                  "confidence": "high", "source_url": "https://...",
                  "rationale": "..."}]}
  commodity_code "0" marks the invoice as a service (not reported).
"""
from __future__ import annotations

import argparse
import copy
import csv
import hashlib
import json
import math
import re
import shutil
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_DOWN
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SOURCE_HEADERS = [
    "Invoice Number", "Count of invoices with current filters", "Invoice Date",
    "Invoice Creation Date", "Commodity weigth", "Commodity code", "Invoice Amount",
    "Invoice Amount in DKK", "Invoice Amount Paid", "Invoice Currency", "Supplier",
    "Vat Code", "Tax Registration Number", "Supplier Type", "Approval Status",
    "Accounting Status", "Conversion Rate", "Accounting Period Name", "Invoice ID",
    "Site Address Country"]
ALIASES = {"Supplier": ["Supplier", "Supplier Name"],
           "Site Address Country": ["Site Address Country", "Taxpayer Country"]}
IND_HEADERS = ["Fakturanummer", "VarekodeRapport", "Varekode", "Vægt",
               "Supplerende enhed", "Mængde supp. enhed", "Service", "Leverandør",
               "Landekode", "Fakturaværdi", "KodeKontrol", "SupplKontrol"]
OUT_HEADERS = ["Varekode", "Transaktionsart", "Landekode", "Vægt",
               "Mængde supp. enhed", "Fakturaværdi"]
TRANS_TYPE = 11
DEFAULT_WEIGHT = Decimal("1")
DEFAULT_SUPPLEMENTARY_QUANTITY = Decimal("1")
INVOICE_DETAIL_FIELDS = [
    ("Supplier Name", ("supplier name",)),
    ("Leverandørvare", ("leverandørvare",)),
    ("Beskrivelse af fakturalinje", ("beskrivelse af fakturalinje",)),
    ("Beskrivelse af fakturakontering", ("beskrivelse af fakturakontering",)),
    ("Mængde", ("mængde", "quantity")),
    ("Forbrugsbasis", ("forbrugsbasis",)),
    ("Type af fakturalinje", ("type af fakturalinje",)),
    ("Kategorinavn", ("kategorinavn",)),
    ("UNSPSC 4 niveauer", ("unspsc 4 niveauer (custom category 1 description)", "unspsc 4 niveauer")),
    ("Udbudskat. 3 niveauer", ("udbudskat. 3 niveauer (custom category 2 description)", "udbudskat. 3 niveauer")),
]
CONFIDENCE_LEVELS = ("confirmed", "high", "medium", "low")
NO_CODE_REASON = "No valid exact match from source, override or LEV-HSKoder"
OUTPUT_NAME_RE = re.compile(r"^(intrastat[ _]((review|log|run)[ _])?\d{4}-\d{2}|lev-hskoder[ _]\d{4}-\d{2})$", re.I)


# --------------------------------------------------------------------------- helpers
def txt(v):
    if v is None or isinstance(v, float) and math.isnan(v):
        return ""
    return str(v).strip()


def hkey(v):
    return re.sub(r"\s+", " ", txt(v)).casefold()


skey = hkey


def num(v):
    if v in (None, ""):
        return None
    if isinstance(v, Decimal):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return Decimal(str(v))
    s = txt(v).replace("\u00a0", "")
    if "," in s and "." not in s:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def excel_num(v):
    if v is None:
        return None
    return int(v) if v == v.to_integral_value() else float(v)


def code(v):
    s = txt(v)
    if re.fullmatch(r"\d+\.0", s):
        s = s[:-2]
    s = re.sub(r"[.\s]", "", s)
    if s.casefold() == "x":
        return "x"
    if not s.isdigit():
        return s
    if s == "0":
        return "0"
    if len(s) < 8:
        s += "00"
    return s[:8]


def country(v):
    s = re.sub(r"\s+", "", txt(v)).upper()
    return s[:2] if re.match(r"^[A-Z]{2}", s) else ""


def get_sheet(wb, name, alts=()):
    for n in (name, *alts):
        if n in wb.sheetnames:
            return wb[n]
    raise ValueError(f"Missing worksheet {name}. Found {wb.sheetnames}")


def first_row(ws):
    for values in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        return values
    return ()


def header_map(values):
    return {hkey(v): i for i, v in enumerate(values, 1) if txt(v)}


def json_value(value):
    if isinstance(value, Decimal):
        return excel_num(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


# --------------------------------------------------------------------------- readers
def source_columns(hm):
    cols = {}
    for canonical in SOURCE_HEADERS:
        col = next((hm[hkey(x)] for x in ALIASES.get(canonical, [canonical]) if hkey(x) in hm), None)
        if col is None:
            return None
        cols[canonical] = col
    return cols


def find_source_sheet(wb):
    ordered = [wb["Rapportdata"]] if "Rapportdata" in wb.sheetnames else []
    ordered += [w for w in wb.worksheets if w not in ordered]
    for ws in ordered:
        cols = source_columns(header_map(first_row(ws)))
        if cols:
            return ws, cols
    return None, None


def source_matches_schema(path):
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return False
    try:
        return find_source_sheet(wb)[0] is not None
    finally:
        wb.close()


def read_source(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws, cols = find_source_sheet(wb)
        if ws is None:
            raise ValueError(f"Could not find Fusion Intrastat source columns in {path}")
        out = []
        for r, values in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            x = {h: (values[c - 1] if c - 1 < len(values) else None) for h, c in cols.items()}
            if any(v not in (None, "") for v in x.values()):
                x["_row"] = r
                out.append(x)
        title = ws.title
    finally:
        wb.close()
    out.sort(key=lambda x: (skey(x["Supplier"]), x["_row"]))
    return out, title


def read_opslag(path):
    """Return (valid codes, unit per code, description per code) from Opslag."""
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = get_sheet(wb, "Opslag")
        valid, units, texts = set(), {}, {}
        for values in ws.iter_rows(min_row=2, max_col=6, values_only=True):
            values = tuple(values) + (None,) * (6 - len(values))
            c = code(values[0])
            if re.fullmatch(r"\d{8}", c):
                valid.add(c)
                if not units.get(c):
                    units[c] = txt(values[1])
                extra = " | ".join(txt(v) for v in values[2:] if txt(v))
                if extra and not texts.get(c):
                    texts[c] = extra
    finally:
        wb.close()
    return valid, units, texts


def read_history(path):
    if not path:
        return {}
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = get_sheet(wb, "LEV-HSKoder", ("LEV -> HS",))
        opts = defaultdict(list)
        for values in ws.iter_rows(min_row=3, max_col=3, values_only=True):
            values = tuple(values) + (None,) * (3 - len(values))
            s, c, n = txt(values[0]), code(values[1]), num(values[2]) or 0
            if s and re.fullmatch(r"\d{8}", c):
                opts[skey(s)].append((c, int(n)))
    finally:
        wb.close()
    return dict(opts)


def read_invoice_details(path):
    if not path:
        return {}, {}
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.active
        header_row = invoice_col = field_cols = None
        # read-only sheets often report max_row=None, so scan a fixed window
        for r, values in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
            hm = {hkey(v): i for i, v in enumerate(values) if txt(v)}
            invoice_col = next((hm[k] for k in ("fakturanr.", "fakturanummer", "invoice number") if k in hm), None)
            field_cols = {name: next((hm[k] for k in aliases if k in hm), None) for name, aliases in INVOICE_DETAIL_FIELDS}
            if invoice_col is not None and all(c is not None for c in field_cols.values()):
                header_row = r
                break
        if header_row is None:
            raise ValueError(f"Could not find invoice number and required invoice-detail columns in {path}")
        quantities, details = defaultdict(Decimal), defaultdict(list)
        for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
            invoice = txt(values[invoice_col] if invoice_col < len(values) else None)
            if not invoice:
                continue
            detail = {name: (values[col] if col < len(values) else None) for name, col in field_cols.items()}
            if any(value not in (None, "") for value in detail.values()):
                details[invoice].append(detail)
            quantity = num(detail["Mængde"])
            if quantity is not None:
                quantities[invoice] += quantity
        return dict(quantities), dict(details)
    finally:
        wb.close()


def read_overrides(path):
    if not path:
        return {}
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))
    else:
        wb = load_workbook(path, data_only=True, read_only=True)
        try:
            it = wb.active.iter_rows(values_only=True)
            hs = [txt(v) for v in next(it, ())]
            rows = [dict(zip(hs, values)) for values in it]
        finally:
            wb.close()
    out = {}
    for x in rows:
        inv = txt(x.get("Invoice Number") or x.get("Fakturanummer"))
        if inv:
            out[inv] = x
    return out


def bool_override(v):
    s = txt(v).casefold()
    if s in {"x", "1", "true", "yes", "ja", "service"}:
        return True
    if s in {"0", "false", "no", "nej", "goods", "vare"}:
        return False
    return None


# --------------------------------------------------------------------------- rules
def process(raw, valid, units, hist_opts, overrides, supplementary_quantities=None,
            invoice_details=None, web=None, pending=None):
    supplementary_quantities = supplementary_quantities or {}
    invoice_details = invoice_details or {}
    web, pending = web or {}, pending or {}
    rows = []
    for src in raw:
        inv = txt(src["Invoice Number"])
        ov = overrides.get(inv, {})
        dec = web.get(inv)
        rawcode = txt(src["Commodity code"])
        sourcecode = code(rawcode)
        supplier = txt(src["Supplier"])
        reasons, selected, origin = [], "", ""

        sv = bool_override(ov.get("Service"))
        service = (sourcecode in {"0", "x"}) if sv is None else sv
        value = num(src["Invoice Amount in DKK"])
        if value is not None and value < 0:
            service = True
            reasons.append("Negative invoice excluded")

        override_code = code(ov.get("Varekode"))
        if service:
            selected, origin = "0", "service"
        elif override_code:
            selected, origin = override_code, "override"
        elif sourcecode in valid:
            selected, origin = sourcecode, "source"
        else:
            candidates = [(c, n) for c, n in hist_opts.get(skey(supplier), []) if c in valid]
            if candidates:
                highest = max(n for _, n in candidates)
                selected = next(c for c, n in candidates if n == highest)
                origin = "supplier history"
        if dec:
            origin = f"web ({dec['confidence']})" + (" service" if service else "")

        if not service and not selected:
            reasons.append(NO_CODE_REASON)
        elif not service and selected not in valid:
            reasons.append(f"Code {selected} is not in Opslag")
        if inv in pending:
            p = pending[inv]
            reasons.append(f"Web suggestion {p['code']} ({p['confidence']}) awaiting confirmation")

        weight = num(ov.get("Vægt") or ov.get("Weight"))
        if weight is None:
            weight = num(src["Commodity weigth"])
        if weight in (None, Decimal("0")):
            weight = DEFAULT_WEIGHT
        land = txt(ov.get("Landekode") or ov.get("Country code")).upper() or country(src["Tax Registration Number"])
        unit = units.get(selected, "") if selected != "0" else ""
        qty = num(ov.get("Mængde supp. enhed") or ov.get("Supplementary quantity"))
        qty_defaulted = False
        if unit and qty is None and not service:
            qty = supplementary_quantities.get(inv)
            if qty is None:
                qty, qty_defaulted = DEFAULT_SUPPLEMENTARY_QUANTITY, True
        if not re.fullmatch(r"[A-Z]{2}", land):
            reasons.append("Invalid country code")
        if value is None:
            reasons.append("Missing DKK value")
        if not supplier:
            reasons.append("Missing supplier")

        base_ok = (not service and selected in valid and re.fullmatch(r"[A-Z]{2}", land)
                   and (not unit or qty is not None) and bool(supplier))
        ready = base_ok and value is not None and value > 0
        excluded = base_ok and value == 0 and not reasons
        status = "SERVICE" if service else "READY" if ready else "EXCLUDED" if excluded else "REVIEW"
        rows.append(dict(
            source=src, source_row=src["_row"], invoice=inv, supplier=supplier, raw_code=rawcode,
            commodity_code=selected, code_origin=origin, weight=weight, unit=unit, quantity=qty,
            qty_defaulted=qty_defaulted, service="x" if service else "", country=land, value=value,
            status=status, reasons=reasons, code_needed=(not service and selected not in valid),
            history=hist_opts.get(skey(supplier), []), invoice_details=invoice_details.get(inv, []),
            web=dec or pending.get(inv), web_applied=bool(dec)))
    return rows


# --------------------------------------------------------------------------- writers
def clear(ws, minrow, maxrow, mincol, maxcol):
    for (r, c), cell in list(ws._cells.items()):
        if minrow <= r <= maxrow and mincol <= c <= maxcol and not isinstance(cell, MergedCell):
            if r == minrow:
                cell.value = None
            else:
                del ws._cells[(r, c)]


def rowstyle(ws, src, dst, c1, c2):
    if src == dst:
        return
    for c in range(c1, c2 + 1):
        a, b = ws.cell(src, c), ws.cell(dst, c)
        if a.has_style:
            b._style = copy.copy(a._style)
        b.number_format = a.number_format
        b.alignment = copy.copy(a.alignment)
        b.protection = copy.copy(a.protection)


def puttext(cell, v):
    cell.number_format = "@"
    cell.value = txt(v)


def setheaders(ws, values, row=1, start=1):
    for c, v in enumerate(values, start):
        ws.cell(row, c).value = v


def style(ws):
    fill = PatternFill("solid", fgColor="1F4E78")
    for c in ws[1]:
        c.font = Font(color="FFFFFF", bold=True)
        c.fill = fill
        c.alignment = Alignment(wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for c in range(1, ws.max_column + 1):
        width = max([len(txt(ws.cell(r, c).value)) for r in range(1, min(ws.max_row, 200) + 1)] or [10]) + 2
        ws.column_dimensions[get_column_letter(c)].width = min(max(width, 10), 45)


def web_fields(x):
    w = x["web"] or {}
    return [w.get("code"), w.get("confidence"), "Ja" if x["web_applied"] else ("Nej" if w else None),
            w.get("source_url"), w.get("rationale")]


def review_record(x):
    w = x["web"]
    return {
        "invoice_number": x["invoice"],
        "status": x["status"],
        "code_needed": x["code_needed"],
        "source_row": x["source_row"],
        "supplier": x["supplier"],
        "source_commodity_code": x["raw_code"] or None,
        "chosen_commodity_code": x["commodity_code"] or None,
        "code_source": x["code_origin"] or None,
        "country_code": x["country"],
        "weight": json_value(x["weight"]),
        "supplementary_unit": x["unit"] or None,
        "supplementary_quantity": json_value(x["quantity"]),
        "invoice_value_dkk": json_value(x["value"]),
        "reasons": x["reasons"],
        "historical_candidates": [{"code": c, "count": n} for c, n in x["history"]],
        "web_suggestion": dict(w, applied=x["web_applied"]) if w else None,
        "invoice_details": [
            {name: json_value(value) if value not in (None, "") else None for name, value in detail.items()}
            for detail in x["invoice_details"]
        ],
    }


def write_review(path, rows, period, step):
    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    setheaders(ws, ["Status", "Code needed", "Source row", "Invoice Number", "Supplier", "Raw code", "Chosen code",
                    "Code source", "Country", "Weight", "Supplementary unit", "Supplementary quantity",
                    "Invoice value DKK", "Reason", "Historical candidates", "Web code", "Web confidence",
                    "Web applied", "Web source", "Web rationale",
                    *[name for name, _ in INVOICE_DETAIL_FIELDS], "Invoice detail line"])
    review = [x for x in rows if x["status"] == "REVIEW"]
    for x in review:
        common = [x["status"], "Ja" if x["code_needed"] else "Nej", x["source_row"], x["invoice"], x["supplier"],
                  x["raw_code"], x["commodity_code"], x["code_origin"], x["country"], excel_num(x["weight"]),
                  x["unit"], excel_num(x["quantity"]), excel_num(x["value"]), "; ".join(x["reasons"]),
                  "; ".join(f"{c} ({n})" for c, n in x["history"]), *web_fields(x)]
        details = x["invoice_details"] or [{}]
        for index, detail in enumerate(details, 1):
            ws.append([*common, *[json_value(detail.get(name)) for name, _ in INVOICE_DETAIL_FIELDS],
                       index if x["invoice_details"] else None])
    style(ws)

    applied = [x for x in rows if x["web_applied"]]
    if applied:
        wa = wb.create_sheet("Web codes applied")
        setheaders(wa, ["Invoice Number", "Supplier", "Status", "Code", "Confidence", "Source URL", "Rationale"])
        for x in applied:
            w = x["web"]
            wa.append([x["invoice"], x["supplier"], x["status"], w["code"], w["confidence"], w["source_url"], w["rationale"]])
        style(wa)
    wb.save(path)

    json_path = path.with_suffix(".json")
    payload = {
        "period": period,
        "step": step,
        "review_count": len(review),
        "code_needed_count": sum(x["code_needed"] for x in review),
        "items": [review_record(x) for x in review],
        "web_codes_applied": [review_record(x) for x in applied],
    }
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return json_path


def build(reference, destination, raw, rows, valid):
    shutil.copy2(reference, destination)
    wb = load_workbook(destination, keep_vba=True, data_only=False)
    rapport = get_sheet(wb, "Rapportdata")
    ind = get_sheet(wb, "Inddata")
    out = get_sheet(wb, "Uddata")
    pivot = get_sheet(wb, "Pivot")
    summ = get_sheet(wb, "Summering")

    clear(rapport, 2, max(rapport.max_row, len(raw) + 1), 1, 20)
    setheaders(rapport, SOURCE_HEADERS)
    text_cols = {"Invoice Number", "Commodity code", "Tax Registration Number", "Invoice ID"}
    for r, x in enumerate(raw, 2):
        rowstyle(rapport, 2, r, 1, 20)
        for c, h in enumerate(SOURCE_HEADERS, 1):
            if h in text_cols:
                puttext(rapport.cell(r, c), x[h])
            else:
                rapport.cell(r, c).value = x[h]

    clear(ind, 2, max(ind.max_row, len(rows) + 1), 1, 20)
    setheaders(ind, IND_HEADERS)
    setheaders(ind, ["Varekode", "Transaktionsart", "Landekode", "Nettovægt", "Supplerende enhed", "Fakturaværdi"], 1, 15)
    for r, x in enumerate(rows, 2):
        rowstyle(ind, 2, r, 1, 20)
        kc = "Service" if x["service"] else "OK" if x["commodity_code"] in valid else "Fejl"
        sc = "Husk mængde" if x["unit"] and x["quantity"] is None and not x["service"] else ""
        vals = [x["invoice"], x["raw_code"], x["commodity_code"], excel_num(x["weight"]), x["unit"],
                excel_num(x["quantity"]), x["service"], x["supplier"], x["country"], excel_num(x["value"]), kc, sc]
        for c, v in enumerate(vals, 1):
            if c in {1, 2, 3, 5, 7, 8, 9, 11, 12}:
                puttext(ind.cell(r, c), v)
            else:
                ind.cell(r, c).value = v
        helper = [x["commodity_code"], TRANS_TYPE, x["country"], excel_num(x["weight"]),
                  excel_num(x["quantity"]), excel_num(x["value"])]
        for c, v in enumerate(helper, 15):
            if c in {15, 17}:
                puttext(ind.cell(r, c), v)
            else:
                ind.cell(r, c).value = v

    ready = [x for x in rows if x["status"] == "READY"]
    clear(pivot, 1, max(pivot.max_row, len(ready) + 5), 1, 7)
    pivot.cell(2, 1).value = "Service"
    pivot.cell(2, 2).value = "(Alle undtagen x)"
    setheaders(pivot, ["Varekode", "Landekode", "Vægt", "Mængde supp. enhed", "Fakturaværdi", "Fakturanummer", "Total"], 5)
    for r, x in enumerate(ready, 6):
        vals = [x["commodity_code"], x["country"], None if x["unit"] else excel_num(x["weight"]),
                excel_num(x["quantity"]) if x["unit"] else "(blank)", excel_num(x["value"]), x["invoice"], None]
        for c, v in enumerate(vals, 1):
            pivot.cell(r, c).value = v

    clear(out, 1, max(out.max_row, len(ready) + 1), 1, 9)
    setheaders(out, OUT_HEADERS)
    out.cell(1, 9).value = TRANS_TYPE
    for r, x in enumerate(ready, 2):
        v = x["value"].quantize(Decimal("1"), rounding=ROUND_DOWN)
        vals = [x["commodity_code"], TRANS_TYPE, x["country"], None if x["unit"] else excel_num(x["weight"]),
                excel_num(x["quantity"]) if x["unit"] else None, excel_num(v)]
        for c, z in enumerate(vals, 1):
            if c in {1, 3}:
                puttext(out.cell(r, c), z)
            else:
                out.cell(r, c).value = z

    groups = defaultdict(lambda: {"w": Decimal(0), "q": Decimal(0), "v": Decimal(0), "n": 0})
    for x in ready:
        g = groups[(x["commodity_code"], x["country"])]
        g["w"] += Decimal(0) if x["unit"] else x["weight"]
        g["q"] += x["quantity"] or Decimal(0)
        g["v"] += x["value"]
        g["n"] += 1
    clear(summ, 1, max(summ.max_row, len(groups) + 1), 1, 19)
    setheaders(summ, ["Varekode", "Transaktionsart", "Landekode", "Nettovægt", "Supplerende enhed", "Fakturaværdi", "Antal fakturaer"])
    for r, ((c, l), g) in enumerate(sorted(groups.items()), 2):
        vals = [c, TRANS_TYPE, l, excel_num(g["w"]) or None, excel_num(g["q"]) or None, excel_num(g["v"]), g["n"]]
        for col, z in enumerate(vals, 1):
            if col in {1, 3}:
                puttext(summ.cell(r, col), z)
            else:
                summ.cell(r, col).value = z

    for name in ("PythonReview", "PythonWebCodes"):
        if name in wb.sheetnames:
            del wb[name]
    rev = wb.create_sheet("PythonReview")
    setheaders(rev, ["Status", "Invoice Number", "Supplier", "Raw code", "Chosen code", "Reason",
                     "Web code", "Web confidence", "Web applied", "Web source", "Web rationale"])
    for x in rows:
        if x["status"] == "REVIEW":
            rev.append([x["status"], x["invoice"], x["supplier"], x["raw_code"], x["commodity_code"],
                        "; ".join(x["reasons"]), *web_fields(x)])
    style(rev)
    web_rows = [x for x in rows if x["web"]]
    if web_rows:
        wc = wb.create_sheet("PythonWebCodes")
        setheaders(wc, ["Invoice Number", "Supplier", "Status", "Code", "Confidence", "Applied", "Source URL", "Rationale"])
        for x in web_rows:
            code_, conf, applied, url, why = web_fields(x)
            wc.append([x["invoice"], x["supplier"], x["status"], code_, conf, applied, url, why])
        style(wc)
    wb.save(destination)
    return len(ready), len(groups)


def last_used_row(ws, max_col=3):
    for r in range(ws.max_row, 1, -1):
        if any(ws.cell(r, c).value not in (None, "") for c in range(1, max_col + 1)):
            return r
    return 1


def delete_period_rows(ws, period):
    """Remove earlier rows for this period so re-runs replace instead of double-count."""
    hits = [r for r in range(2, ws.max_row + 1) if txt(ws.cell(r, 1).value) == period]
    i = len(hits) - 1
    while i >= 0:
        end = start = hits[i]
        while i > 0 and hits[i - 1] == start - 1:
            i -= 1
            start = hits[i]
        ws.delete_rows(start, end - start + 1)
        i -= 1
    return len(hits)


def update_hs(src, dst, period, rows):
    if Path(src).resolve() != Path(dst).resolve():
        shutil.copy2(src, dst)
    wb = load_workbook(dst, data_only=False)
    sam = get_sheet(wb, "Samlet")
    eligible = [x for x in rows if x["status"] in {"READY", "SERVICE"} and x["supplier"] and x["commodity_code"]]
    signature = hashlib.sha256((period + "\n" + "\n".join(
        "|".join([x["invoice"], x["supplier"], x["commodity_code"], txt(x["value"])]) for x in eligible)).encode()).hexdigest()

    replaced = delete_period_rows(sam, period)
    r = last_used_row(sam) + 1
    for x in eligible:
        sam.cell(r, 1).value = period
        puttext(sam.cell(r, 2), x["commodity_code"])
        sam.cell(r, 3).value = x["supplier"]
        r += 1

    log = wb["PythonRunLog"] if "PythonRunLog" in wb.sheetnames else wb.create_sheet("PythonRunLog")
    log.sheet_state = "hidden"
    if not txt(log.cell(1, 1).value):
        log.append(["Timestamp", "Period", "Signature", "Rows", "Replaced rows"])
    log.append([datetime.now().isoformat(timespec="seconds"), period, signature, len(eligible), replaced])

    counts, display = Counter(), {}
    for r in range(2, sam.max_row + 1):
        s, c = txt(sam.cell(r, 3).value), code(sam.cell(r, 2).value)
        if s and c:
            k = skey(s)
            display.setdefault(k, s)
            counts[(k, c)] += 1

    sm = get_sheet(wb, "LEV-HSKoder", ("LEV -> HS",))
    clear(sm, 3, max(sm.max_row, 3), 1, 3)
    setheaders(sm, ["Leverandør", "Varekode", "Antal af Varekode"], 2)
    bys = defaultdict(list)
    r = 3
    for (k, c), n in sorted(counts.items(), key=lambda z: (display[z[0][0]].casefold(), z[0][1])):
        sm.cell(r, 1).value = display[k]
        puttext(sm.cell(r, 2), c)
        sm.cell(r, 3).value = n
        bys[k].append((c, n))
        r += 1

    rec = get_sheet(wb, "HS KODE")
    clear(rec, 3, max(rec.max_row, 3), 1, 2)
    setheaders(rec, ["Leverandør", "Hs kode"], 2)
    r = 3
    for k in sorted(bys, key=lambda x: display[x].casefold()):
        opts = [(c, n) for c, n in bys[k] if re.fullmatch(r"\d{8}", c)]
        val = "N/A"
        if opts:
            m = max(n for _, n in opts)
            top = sorted({c for c, n in opts if n == m})
            val = top[0] if len(top) == 1 else "TIE: " + ", ".join(top)
        rec.cell(r, 1).value = display[k]
        puttext(rec.cell(r, 2), val)
        r += 1
    wb.save(dst)
    return {"appended": len(eligible), "replaced_existing_period_rows": replaced}


# --------------------------------------------------------------------------- decisions (step 2)
def load_decisions(path):
    data = json.loads(Path(path).read_text(encoding="utf-8-sig"))
    if isinstance(data, list):
        return None, data
    return data.get("period"), data.get("decisions", [])


def validate_decisions(items, needs_code, valid, include_low):
    web, pending, rejected = {}, {}, []
    for d in items:
        inv = txt(d.get("invoice_number") or d.get("Invoice Number") or d.get("Fakturanummer"))
        c = code(d.get("commodity_code") or d.get("Varekode"))
        if c == "x":
            c = "0"
        conf = txt(d.get("confidence")).casefold() or "low"
        entry = dict(invoice=inv, code=c, confidence=conf, source_url=txt(d.get("source_url")),
                     rationale=txt(d.get("rationale")))
        problem = None
        if not inv:
            problem = "missing invoice_number"
        elif inv not in needs_code:
            problem = "invoice is not a REVIEW row that needs a commodity code"
        elif inv in web or inv in pending:
            problem = "duplicate decision for this invoice"
        elif conf not in CONFIDENCE_LEVELS:
            problem = f"confidence must be one of {', '.join(CONFIDENCE_LEVELS)}"
        elif c != "0" and c not in valid:
            problem = f"code {c or '(empty)'} is not a valid 8-digit code in Opslag"
        if problem:
            rejected.append(dict(entry, problem=problem))
            continue
        hold = conf == "low" or (c == "0" and conf == "medium")
        if hold and not include_low:
            pending[inv] = entry
        else:
            web[inv] = entry
    return web, pending, rejected


# --------------------------------------------------------------------------- discovery
def default_data_dir():
    here = Path(__file__).resolve().parent if "__file__" in globals() else Path.cwd()
    return here / "data" if (here / "data").is_dir() else here


def excel_files(data_dir):
    if not data_dir.is_dir():
        return []
    return sorted((p for p in data_dir.iterdir()
                   if p.is_file() and p.suffix.lower() in {".xlsx", ".xlsm"}
                   and not p.name.startswith("~$") and not OUTPUT_NAME_RE.match(p.stem)),
                  key=lambda x: x.name.casefold())


def is_invoice_file(p):
    lower = p.name.casefold()
    return "fakturaer" in lower or "invoice" in lower


def discover_default_files(data_dir):
    files = [p for p in excel_files(data_dir) if not is_invoice_file(p)]
    ref = next((p for p in files if "skabelon" in p.name.casefold() and p.suffix.lower() == ".xlsm"), None)
    hs = next((p for p in files if "hskoder" in p.name.casefold()), None)
    source = next((p for p in files if p not in (ref, hs) and source_matches_schema(p)), None)
    return {"source": source, "reference": ref, "hskoder": hs}


def discover_invoice_details(data_dir, period):
    matches = [p for p in excel_files(data_dir) if is_invoice_file(p) and infer_period(p) == period]
    if len(matches) > 1:
        raise ValueError(f"Multiple invoice-details workbooks found for {period}: {matches}")
    return matches[0] if matches else None


MONTHS = {"januar": "01", "january": "01", "februar": "02", "february": "02", "marts": "03", "march": "03",
          "april": "04", "maj": "05", "may": "05", "juni": "06", "june": "06", "juli": "07", "july": "07",
          "august": "08", "september": "09", "oktober": "10", "october": "10", "november": "11",
          "december": "12"}


def infer_period(path):
    stem = path.stem
    m = re.search(r"(\d{4})-(0[1-9]|1[0-2])", stem)
    if m:
        return f"{m.group(1)}-{m.group(2)}"
    m = re.search(r"(\d{4})", stem)
    if m:
        for name, month in MONTHS.items():
            if name in stem.casefold():
                return f"{m.group(1)}-{month}"
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return None
    try:
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(min_row=1, max_row=200, values_only=True))
            if not rows:
                continue
            hdrs = [hkey(v) for v in rows[0]]
            for row in rows[1:]:
                data = {hdrs[i]: row[i] for i in range(min(len(hdrs), len(row))) if hdrs[i]}
                invoice_date = data.get("invoice date")
                period_value = data.get("accounting period name")
                if period_value is None and invoice_date is not None:
                    period_value = invoice_date
                if not period_value:
                    continue
                if hasattr(period_value, "year") and hasattr(period_value, "month"):
                    return f"{period_value.year:04d}-{period_value.month:02d}"
                if isinstance(period_value, str):
                    m = re.search(r"(\d{4})[-.](0[1-9]|1[0-2])", period_value)
                    if m:
                        return f"{m.group(1)}-{m.group(2)}"
                    m = re.search(r"(\d{2})-(\d{2})", period_value)
                    if m:
                        year = invoice_date.year if hasattr(invoice_date, "year") else datetime.now().year
                        return f"{year:04d}-{m.group(1)}"
        return None
    finally:
        wb.close()


# --------------------------------------------------------------------------- orchestration
def resolve_inputs(a):
    data_dir = Path(a.data_dir) if a.data_dir else default_data_dir()
    defaults = discover_default_files(data_dir)
    source = a.source or defaults["source"]
    reference = a.reference or defaults["reference"]
    hskoder = a.hskoder or defaults["hskoder"]
    if not source:
        raise SystemExit(f"No valid Fusion Intrastat source workbook was found in {data_dir}.")
    if not reference:
        raise SystemExit(f"No reference workbook (name containing 'skabelon', .xlsm) was found in {data_dir}.")
    period = a.period or infer_period(Path(source))
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", period or ""):
        raise SystemExit("Period must be YYYY-MM (could not infer it; pass --period).")
    invoice_details = a.invoice_details or discover_invoice_details(data_dir, period)
    out = Path(a.output_dir) if a.output_dir else data_dir / "output"
    absolute = lambda p: str(Path(p).resolve()) if p else None
    return {"period": period, "source": absolute(source), "reference": absolute(reference),
            "hskoder": absolute(hskoder), "invoice_details": absolute(invoice_details),
            "overrides": absolute(a.overrides), "output_dir": absolute(out)}


def output_paths(inputs):
    out, period = Path(inputs["output_dir"]), inputs["period"]
    # No spaces: Copilot only turns simple file names into download links.
    return {"intrastat": out / f"Intrastat_{period}.xlsm",
            "review": out / f"Intrastat_review_{period}.xlsx",
            "hskoder": out / f"LEV-HSKoder_{period}.xlsx" if inputs["hskoder"] else None,
            "log": out / f"Intrastat_log_{period}.json",
            "manifest": out / f"Intrastat_run_{period}.json"}


def run_pipeline(inputs, decisions=None, include_low=True):
    period = inputs["period"]
    paths = output_paths(inputs)
    Path(inputs["output_dir"]).mkdir(parents=True, exist_ok=True)

    raw, source_sheet = read_source(inputs["source"])
    valid, units, _ = read_opslag(inputs["reference"])
    hist = read_history(inputs["hskoder"])
    ovs = read_overrides(inputs["overrides"])
    quantities, details = read_invoice_details(inputs["invoice_details"])

    web, pending, rejected = {}, {}, []
    step = "run"
    if decisions is not None:
        step = "apply"
        base = process(raw, valid, units, hist, ovs, quantities, details)
        needs_code = {x["invoice"] for x in base if x["status"] == "REVIEW" and x["code_needed"]}
        web, pending, rejected = validate_decisions(decisions, needs_code, valid, include_low)
        ovs = dict(ovs)
        for inv, d in web.items():
            o = dict(ovs.get(inv, {}))
            if d["code"] == "0":
                o["Service"] = "x"
            else:
                o["Varekode"] = d["code"]
            ovs[inv] = o

    rows = process(raw, valid, units, hist, ovs, quantities, details, web, pending)
    ready, groups = build(Path(inputs["reference"]), paths["intrastat"], raw, rows, valid)
    review_json = write_review(paths["review"], rows, period, step)
    hs_result = update_hs(inputs["hskoder"], paths["hskoder"], period, rows) if inputs["hskoder"] else None

    counts = Counter(x["status"] for x in rows)
    log = {
        "step": step,
        "period": period,
        "source_sheet": source_sheet,
        "inputs": inputs,
        "counts": {"source": len(rows), "ready": counts["READY"], "service": counts["SERVICE"],
                   "excluded": counts["EXCLUDED"], "review": counts["REVIEW"],
                   "review_code_needed": sum(x["status"] == "REVIEW" and x["code_needed"] for x in rows),
                   "uddata": ready, "summering_groups": groups,
                   "supplementary_quantity_defaulted_to_1": sum(x["qty_defaulted"] for x in rows if x["status"] == "READY")},
        "hskoder": hs_result,
        "outputs": {"intrastat": str(paths["intrastat"]), "review": str(paths["review"]),
                    "review_json": str(review_json), "hskoder": str(paths["hskoder"]) if paths["hskoder"] else None,
                    "manifest": str(paths["manifest"])},
    }
    if step == "apply":
        log["web"] = {"applied": len(web),
                      "applied_by_confidence": dict(Counter(d["confidence"] for d in web.values())),
                      "pending_confirmation": list(pending.values()),
                      "rejected": rejected}
    else:
        paths["manifest"].write_text(json.dumps(dict(inputs, created=datetime.now().isoformat(timespec="seconds")),
                                                ensure_ascii=False, indent=2), encoding="utf-8")
    paths["log"].write_text(json.dumps(log, ensure_ascii=False, indent=2), encoding="utf-8")
    return log


def find_manifest(a, period):
    if a.manifest:
        return Path(a.manifest)
    data_dir = Path(a.data_dir) if a.data_dir else default_data_dir()
    out = Path(a.output_dir) if a.output_dir else data_dir / "output"
    if period:
        return out / f"Intrastat_run_{period}.json"
    found = sorted(out.glob("Intrastat_run_*.json"))
    if len(found) != 1:
        raise SystemExit("Pass --manifest (or a decisions file with 'period'); found "
                         f"{len(found)} run manifests in {out}.")
    return found[0]


def cmd_run(a):
    return run_pipeline(resolve_inputs(a))


def cmd_apply(a):
    period, items = load_decisions(a.decisions)
    manifest = find_manifest(a, period)
    if not manifest.is_file():
        raise SystemExit(f"Run manifest {manifest} not found. Run step 1 ('run') first.")
    inputs = json.loads(manifest.read_text(encoding="utf-8"))
    inputs.pop("created", None)
    if period and period != inputs["period"]:
        raise SystemExit(f"Decisions period {period} does not match manifest period {inputs['period']}.")
    return run_pipeline(inputs, decisions=items, include_low=not a.hold_low)


def cmd_codes(a):
    if a.reference:
        reference = a.reference
    elif a.manifest:
        reference = json.loads(Path(a.manifest).read_text(encoding="utf-8"))["reference"]
    else:
        reference = discover_default_files(Path(a.data_dir) if a.data_dir else default_data_dir())["reference"]
    if not reference:
        raise SystemExit("No reference workbook found; pass --reference or --manifest.")
    valid, units, texts = read_opslag(reference)
    prefixes = [re.sub(r"[.\s]", "", p) for p in (a.prefix or [])]
    needle = (a.search or "").casefold()
    hits = [c for c in sorted(valid)
            if (not prefixes or any(c.startswith(p) for p in prefixes))
            and (not needle or needle in texts.get(c, "").casefold())]
    return {"reference": str(reference), "total_matches": len(hits),
            "codes": [{"code": c, "supplementary_unit": units.get(c) or None, "text": texts.get(c) or None}
                      for c in hits[: a.limit]]}


def build_parser():
    p = argparse.ArgumentParser(description="Intrastat pipeline (step 1: run, step 2: apply)")
    sub = p.add_subparsers(dest="command", required=True)

    r = sub.add_parser("run", help="Step 1: build Intrastat, review and LEV-HSKoder from monthly inputs")
    r.add_argument("--data-dir", help="Folder with the input files (default: ./data next to the script, else the script folder)")
    r.add_argument("--source")
    r.add_argument("--reference")
    r.add_argument("--period")
    r.add_argument("--hskoder")
    r.add_argument("--invoice-details")
    r.add_argument("--overrides")
    r.add_argument("--output-dir")
    r.set_defaults(func=cmd_run)

    ap = sub.add_parser("apply", help="Step 2: apply web-found codes and rebuild all outputs")
    ap.add_argument("--decisions", required=True)
    ap.add_argument("--manifest", help="Intrastat_run_YYYY-MM.json from step 1 (default: found in output dir)")
    ap.add_argument("--data-dir")
    ap.add_argument("--output-dir")
    ap.add_argument("--hold-low", action="store_true", help="Keep low-confidence codes in review instead of applying them")
    ap.set_defaults(func=cmd_apply)

    c = sub.add_parser("codes", help="List valid Opslag codes (step 2 helper)")
    c.add_argument("--prefix", action="append", help="HS/CN prefix, e.g. 8471 (repeatable)")
    c.add_argument("--search", help="Text to find in the Opslag description columns")
    c.add_argument("--limit", type=int, default=60)
    c.add_argument("--reference")
    c.add_argument("--manifest")
    c.add_argument("--data-dir")
    c.set_defaults(func=cmd_codes)
    return p


def main(argv=None):
    argv = list(sys.argv[1:] if argv is None else argv)
    if not argv or argv[0].startswith("-"):
        argv.insert(0, "run")
    a = build_parser().parse_args(argv)
    result = a.func(a)
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
    return result


if __name__ == "__main__":
    main()
